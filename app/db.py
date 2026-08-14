import csv
import io
import sqlite3
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "demo.db"

CATEGORIES = ["Laptop", "Server", "PLC", "Industrial PC", "Network Switch", "Sensor", "Gateway"]
STATUSES = ["In Service", "Maintenance", "Retired", "Spare"]
LOCATIONS = ["Factory Alpha", "Workshop A", "Lab B", "Office C"]

SEED_ASSETS = [
    {
        "asset_id": "ITOT-001",
        "asset_name": "Engineering Laptop Alpha",
        "category": "Laptop",
        "location": "Office C",
        "owner": "Iris Cole",
        "status": "In Service",
        "last_inventory_date": "2026-07-20",
        "notes": "Synthetic office engineering device for portfolio demo.",
    },
    {
        "asset_id": "ITOT-002",
        "asset_name": "Virtualization Host Beta",
        "category": "Server",
        "location": "Lab B",
        "owner": "Noah Grant",
        "status": "Maintenance",
        "last_inventory_date": "2026-07-18",
        "notes": "Synthetic lab server used to demonstrate IT asset tracking.",
    },
    {
        "asset_id": "ITOT-003",
        "asset_name": "Packaging PLC Alpha",
        "category": "PLC",
        "location": "Factory Alpha",
        "owner": "Mia Turner",
        "status": "In Service",
        "last_inventory_date": "2026-07-25",
        "notes": "Synthetic OT controller for packaging line showcase.",
    },
    {
        "asset_id": "ITOT-004",
        "asset_name": "Vision Cell IPC 01",
        "category": "Industrial PC",
        "location": "Workshop A",
        "owner": "Luca Reed",
        "status": "Spare",
        "last_inventory_date": "2026-07-11",
        "notes": "Synthetic industrial PC prepared as spare unit.",
    },
    {
        "asset_id": "ITOT-005",
        "asset_name": "Lab Core Switch",
        "category": "Network Switch",
        "location": "Lab B",
        "owner": "Ava Brooks",
        "status": "In Service",
        "last_inventory_date": "2026-07-21",
        "notes": "Synthetic managed switch for demo network topology.",
    },
    {
        "asset_id": "ITOT-006",
        "asset_name": "Temperature Sensor Node 7",
        "category": "Sensor",
        "location": "Factory Alpha",
        "owner": "Ethan Vale",
        "status": "Maintenance",
        "last_inventory_date": "2026-07-15",
        "notes": "Synthetic plant-floor sensing asset for condition monitoring demo.",
    },
    {
        "asset_id": "ITOT-007",
        "asset_name": "Edge Gateway Delta",
        "category": "Gateway",
        "location": "Workshop A",
        "owner": "Sophia Lane",
        "status": "In Service",
        "last_inventory_date": "2026-07-23",
        "notes": "Synthetic gateway bridging IT and OT segments in the demo.",
    },
    {
        "asset_id": "ITOT-008",
        "asset_name": "QA Laptop Bravo",
        "category": "Laptop",
        "location": "Office C",
        "owner": "Oliver Nash",
        "status": "Retired",
        "last_inventory_date": "2026-06-30",
        "notes": "Synthetic retired user endpoint used for lifecycle history display.",
    },
]

SEED_HISTORY = [
    ("ITOT-002", "status_change", "In Service", "Maintenance", "Synthetic firmware review and scheduled maintenance."),
    ("ITOT-004", "maintenance", "In Service", "Spare", "Synthetic spare preparation after workstation refresh."),
    ("ITOT-008", "status_change", "Maintenance", "Retired", "Synthetic decommission scenario for public demo."),
]

SEED_INVENTORY = [
    ("ITOT-001", "2026-07-20", "Matched", "Office C", "Synthetic cycle count completed."),
    ("ITOT-003", "2026-07-25", "Matched", "Factory Alpha", "Synthetic OT cabinet inventory check."),
    ("ITOT-006", "2026-07-15", "Location Verified", "Factory Alpha", "Synthetic maintenance bench verification."),
]


def get_connection() -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with get_connection() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS assets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                asset_id TEXT NOT NULL UNIQUE,
                asset_name TEXT NOT NULL,
                category TEXT NOT NULL,
                location TEXT NOT NULL,
                owner TEXT NOT NULL,
                status TEXT NOT NULL,
                last_inventory_date TEXT,
                notes TEXT NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS asset_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                asset_id INTEGER NOT NULL,
                event_type TEXT NOT NULL,
                from_status TEXT,
                to_status TEXT,
                notes TEXT NOT NULL,
                event_date TEXT NOT NULL DEFAULT CURRENT_DATE,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(asset_id) REFERENCES assets(id)
            );

            CREATE TABLE IF NOT EXISTS inventory_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                asset_id INTEGER NOT NULL,
                inventory_date TEXT NOT NULL,
                outcome TEXT NOT NULL,
                counted_location TEXT NOT NULL,
                notes TEXT NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(asset_id) REFERENCES assets(id)
            );

            CREATE INDEX IF NOT EXISTS idx_assets_category ON assets(category);
            CREATE INDEX IF NOT EXISTS idx_assets_status ON assets(status);
            CREATE INDEX IF NOT EXISTS idx_assets_location ON assets(location);
            CREATE INDEX IF NOT EXISTS idx_history_asset ON asset_history(asset_id);
            CREATE INDEX IF NOT EXISTS idx_inventory_asset ON inventory_records(asset_id);
            """
        )
        seed_data(conn)


def seed_data(conn: sqlite3.Connection) -> None:
    existing = conn.execute("SELECT COUNT(*) AS count FROM assets").fetchone()["count"]
    if existing:
        return
    conn.executemany(
        """
        INSERT INTO assets(asset_id, asset_name, category, location, owner, status, last_inventory_date, notes)
        VALUES(:asset_id, :asset_name, :category, :location, :owner, :status, :last_inventory_date, :notes)
        """,
        SEED_ASSETS,
    )
    asset_map = {row["asset_id"]: row["id"] for row in conn.execute("SELECT id, asset_id FROM assets").fetchall()}
    conn.executemany(
        """
        INSERT INTO asset_history(asset_id, event_type, from_status, to_status, notes)
        VALUES(?, ?, ?, ?, ?)
        """,
        [(asset_map[asset_code], event_type, from_status, to_status, notes) for asset_code, event_type, from_status, to_status, notes in SEED_HISTORY],
    )
    conn.executemany(
        """
        INSERT INTO inventory_records(asset_id, inventory_date, outcome, counted_location, notes)
        VALUES(?, ?, ?, ?, ?)
        """,
        [(asset_map[asset_code], inventory_date, outcome, counted_location, notes) for asset_code, inventory_date, outcome, counted_location, notes in SEED_INVENTORY],
    )


def asset_filters() -> dict[str, list[str]]:
    return {"categories": CATEGORIES, "statuses": STATUSES, "locations": LOCATIONS}


def dashboard_summary() -> dict[str, Any]:
    with get_connection() as conn:
        totals = conn.execute(
            """
            SELECT COUNT(*) AS asset_count,
                   SUM(CASE WHEN status = 'In Service' THEN 1 ELSE 0 END) AS in_service_count,
                   SUM(CASE WHEN status = 'Maintenance' THEN 1 ELSE 0 END) AS maintenance_count,
                   SUM(CASE WHEN category IN ('PLC', 'Industrial PC', 'Sensor', 'Gateway') THEN 1 ELSE 0 END) AS ot_count
            FROM assets
            """
        ).fetchone()
        by_status = conn.execute("SELECT status, COUNT(*) AS count FROM assets GROUP BY status ORDER BY count DESC, status").fetchall()
        by_category = conn.execute("SELECT category, COUNT(*) AS count FROM assets GROUP BY category ORDER BY count DESC, category").fetchall()
        recent_history = conn.execute(
            """
            SELECT h.*, a.asset_id AS asset_code, a.asset_name
            FROM asset_history h
            JOIN assets a ON a.id = h.asset_id
            ORDER BY h.created_at DESC, h.id DESC
            LIMIT 6
            """
        ).fetchall()
        recent_inventory = conn.execute(
            """
            SELECT i.*, a.asset_id AS asset_code, a.asset_name
            FROM inventory_records i
            JOIN assets a ON a.id = i.asset_id
            ORDER BY i.inventory_date DESC, i.id DESC
            LIMIT 6
            """
        ).fetchall()
    return {
        "totals": dict(totals),
        "by_status": [dict(row) for row in by_status],
        "by_category": [dict(row) for row in by_category],
        "recent_history": [dict(row) for row in recent_history],
        "recent_inventory": [dict(row) for row in recent_inventory],
    }


def list_assets(filters: dict[str, str]) -> list[dict[str, Any]]:
    conditions = []
    params: list[Any] = []
    if filters.get("q"):
        conditions.append("(asset_id LIKE ? OR asset_name LIKE ? OR owner LIKE ? OR notes LIKE ?)")
        q = f"%{filters['q'].strip()}%"
        params.extend([q, q, q, q])
    for field in ["category", "status", "location"]:
        if filters.get(field):
            conditions.append(f"{field} = ?")
            params.append(filters[field])
    where = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    with get_connection() as conn:
        rows = conn.execute(
            f"SELECT * FROM assets {where} ORDER BY updated_at DESC, id DESC",
            params,
        ).fetchall()
    return [dict(row) for row in rows]


def get_asset(asset_pk: int) -> dict[str, Any] | None:
    with get_connection() as conn:
        row = conn.execute("SELECT * FROM assets WHERE id = ?", (asset_pk,)).fetchone()
    return dict(row) if row else None


def asset_exists_by_code(asset_code: str, exclude_id: int | None = None) -> bool:
    query = "SELECT id FROM assets WHERE asset_id = ?"
    params: list[Any] = [asset_code]
    if exclude_id:
        query += " AND id != ?"
        params.append(exclude_id)
    with get_connection() as conn:
        return conn.execute(query, params).fetchone() is not None


def create_asset(payload: dict[str, Any]) -> int:
    with get_connection() as conn:
        cursor = conn.execute(
            """
            INSERT INTO assets(asset_id, asset_name, category, location, owner, status, last_inventory_date, notes, updated_at)
            VALUES(:asset_id, :asset_name, :category, :location, :owner, :status, :last_inventory_date, :notes, CURRENT_TIMESTAMP)
            """,
            payload,
        )
        asset_pk = int(cursor.lastrowid)
        conn.execute(
            """
            INSERT INTO asset_history(asset_id, event_type, from_status, to_status, notes, event_date)
            VALUES(?, 'created', NULL, ?, ?, COALESCE(?, CURRENT_DATE))
            """,
            (asset_pk, payload["status"], "Synthetic asset created in the public demo system.", payload.get("last_inventory_date")),
        )
        return asset_pk


def update_asset(asset_pk: int, payload: dict[str, Any], previous_status: str) -> None:
    with get_connection() as conn:
        conn.execute(
            """
            UPDATE assets
            SET asset_id = :asset_id,
                asset_name = :asset_name,
                category = :category,
                location = :location,
                owner = :owner,
                status = :status,
                last_inventory_date = :last_inventory_date,
                notes = :notes,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = :id
            """,
            {**payload, "id": asset_pk},
        )
        if previous_status != payload["status"]:
            conn.execute(
                """
                INSERT INTO asset_history(asset_id, event_type, from_status, to_status, notes)
                VALUES(?, 'status_change', ?, ?, ?)
                """,
                (asset_pk, previous_status, payload["status"], "Status updated through the asset edit form."),
            )


def list_asset_history(asset_pk: int) -> list[dict[str, Any]]:
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT * FROM asset_history WHERE asset_id = ? ORDER BY event_date DESC, id DESC",
            (asset_pk,),
        ).fetchall()
    return [dict(row) for row in rows]


def list_inventory_records(asset_pk: int | None = None) -> list[dict[str, Any]]:
    query = """
        SELECT i.*, a.asset_id AS asset_code, a.asset_name
        FROM inventory_records i
        JOIN assets a ON a.id = i.asset_id
    """
    params: list[Any] = []
    if asset_pk is not None:
        query += " WHERE i.asset_id = ?"
        params.append(asset_pk)
    query += " ORDER BY i.inventory_date DESC, i.id DESC"
    with get_connection() as conn:
        rows = conn.execute(query, params).fetchall()
    return [dict(row) for row in rows]


def create_inventory_record(asset_pk: int, inventory_date: str, outcome: str, counted_location: str, notes: str) -> None:
    with get_connection() as conn:
        conn.execute(
            """
            INSERT INTO inventory_records(asset_id, inventory_date, outcome, counted_location, notes)
            VALUES(?, ?, ?, ?, ?)
            """,
            (asset_pk, inventory_date, outcome, counted_location, notes),
        )
        conn.execute(
            "UPDATE assets SET last_inventory_date = ?, location = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (inventory_date, counted_location, asset_pk),
        )
        conn.execute(
            """
            INSERT INTO asset_history(asset_id, event_type, from_status, to_status, notes, event_date)
            VALUES(?, 'inventory', NULL, NULL, ?, ?)
            """,
            (asset_pk, f"Inventory outcome: {outcome}. {notes}".strip(), inventory_date),
        )


def create_history_record(asset_pk: int, event_type: str, notes: str, from_status: str | None = None, to_status: str | None = None, event_date: str | None = None) -> None:
    with get_connection() as conn:
        conn.execute(
            """
            INSERT INTO asset_history(asset_id, event_type, from_status, to_status, notes, event_date)
            VALUES(?, ?, ?, ?, ?, COALESCE(?, CURRENT_DATE))
            """,
            (asset_pk, event_type, from_status, to_status, notes, event_date),
        )
        conn.execute("UPDATE assets SET updated_at = CURRENT_TIMESTAMP WHERE id = ?", (asset_pk,))


def import_assets(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    rows: list[dict[str, str]]
    if suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    elif suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(cell).strip() for cell in values[0]]
        rows = []
        for values_row in values[1:]:
            rows.append({headers[idx]: "" if idx >= len(values_row) or values_row[idx] is None else str(values_row[idx]).strip() for idx in range(len(headers))})
    else:
        raise ValueError("仅支持导入 CSV 和 XLSX 文件。")

    imported = 0
    skipped: list[str] = []
    required = ["Asset ID", "Asset Name", "Category", "Location", "Owner", "Status", "Last Inventory Date", "Notes"]
    if rows and any(col not in rows[0] for col in required):
        raise ValueError("导入文件缺少一个或多个必填列。")

    for index, row in enumerate(rows, start=2):
        payload = {
            "asset_id": row.get("Asset ID", "").strip(),
            "asset_name": row.get("Asset Name", "").strip(),
            "category": row.get("Category", "").strip(),
            "location": row.get("Location", "").strip(),
            "owner": row.get("Owner", "").strip(),
            "status": row.get("Status", "").strip(),
            "last_inventory_date": row.get("Last Inventory Date", "").strip(),
            "notes": row.get("Notes", "").strip(),
        }
        if not all(payload.values()):
            skipped.append(f"第 {index} 行：存在必填值为空。")
            continue
        if payload["category"] not in CATEGORIES or payload["status"] not in STATUSES or payload["location"] not in LOCATIONS:
            skipped.append(f"第 {index} 行：资产分类、状态或位置不合法。")
            continue
        if asset_exists_by_code(payload["asset_id"]):
            skipped.append(f"第 {index} 行：Asset ID 已存在。")
            continue
        create_asset(payload)
        imported += 1
    return {"imported": imported, "skipped": skipped}


def export_assets_csv() -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Asset ID", "Asset Name", "Category", "Location", "Owner", "Status", "Last Inventory Date", "Notes"])
    for asset in list_assets({}):
        writer.writerow([
            asset["asset_id"],
            asset["asset_name"],
            asset["category"],
            asset["location"],
            asset["owner"],
            asset["status"],
            asset["last_inventory_date"],
            asset["notes"],
        ])
    return output.getvalue().encode("utf-8-sig")


def export_assets_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Assets"
    sheet.append(["Asset ID", "Asset Name", "Category", "Location", "Owner", "Status", "Last Inventory Date", "Notes"])
    for asset in list_assets({}):
        sheet.append([
            asset["asset_id"],
            asset["asset_name"],
            asset["category"],
            asset["location"],
            asset["owner"],
            asset["status"],
            asset["last_inventory_date"],
            asset["notes"],
        ])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def today_iso() -> str:
    return date.today().isoformat()
